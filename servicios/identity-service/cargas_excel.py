"""Lectura y escritura XLSX limitada. No evalúa fórmulas ni ejecuta macros."""
import re
from defusedxml.ElementTree import iterparse
from collections import Counter
from io import BytesIO
from zipfile import BadZipFile, ZipFile
from openpyxl import Workbook, load_workbook, DEFUSEDXML
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from usuarios import IdentidadCrear

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
HEADERS = ("id_externo", "nombre", "email", "rol")
ROLES = ("LIDER_FACTURACION", "LIDER_LOGISTICA", "LIDER_GESTION_HUMANA")

class ExcelInvalido(ValueError):
    pass

def _archivo_seguro(content, filename):
    if not filename.lower().endswith(".xlsx"):
        raise ExcelInvalido("Use la plantilla Excel .xlsx; no se admiten .xls ni .xlsm.")
    if not content or len(content) > MAX_FILE_BYTES:
        raise ExcelInvalido("El archivo debe tener contenido y pesar como máximo 2 MiB.")
    if not DEFUSEDXML:
        raise RuntimeError("El lector XML protegido no está disponible")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > 128 or len(names) != len(set(names)):
                raise ExcelInvalido("Estructura Excel no admitida.")
            if sum(entry.file_size for entry in entries) > 10 * 1024 * 1024:
                raise ExcelInvalido("El Excel expandido excede el límite.")
            for entry in entries:
                if entry.flag_bits & 1 or entry.file_size > 5 * 1024 * 1024:
                    raise ExcelInvalido("Excel cifrado o demasiado grande.")
                if entry.file_size > 100 * max(entry.compress_size, 1):
                    raise ExcelInvalido("Compresión Excel excesiva.")
                if "vbaproject" in entry.filename.lower() or "externallinks/" in entry.filename.lower():
                    raise ExcelInvalido("No se admiten macros ni vínculos externos.")
    except (BadZipFile, OSError) as error:
        raise ExcelInvalido("No es un archivo Excel .xlsx válido.") from error

def leer_usuarios(content: bytes, filename: str, compania_id: int) -> list[dict]:
    _archivo_seguro(content, filename)
    workbook = None
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False,
                                 keep_links=False, keep_vba=False)
        if "Usuarios" not in workbook.sheetnames:
            raise ExcelInvalido("Falta la hoja Usuarios. Descargue la plantilla.")
        sheet = workbook["Usuarios"]
        # Comprobar referencias físicas: no confiar solo en dimension, que puede mentir.
        with ZipFile(BytesIO(content)) as archive:
            for event, element in iterparse(BytesIO(archive.read(sheet._worksheet_path)), events=("start", "end")):
                tag = element.tag.rsplit("}", 1)[-1]
                if event == "start" and tag == "row":
                    if not 1 <= int(element.attrib.get("r", "0")) <= MAX_ROWS + 1:
                        raise ExcelInvalido("Use como máximo 500 filas de usuarios.")
                if event == "start" and tag == "c":
                    match = re.fullmatch(r"[A-D]([1-9][0-9]*)", element.attrib.get("r", ""))
                    if not match or int(match.group(1)) > MAX_ROWS + 1:
                        raise ExcelInvalido("Use cuatro columnas y como máximo 500 filas de usuarios.")
                if event == "end":
                    element.clear()
        if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > len(HEADERS):
            raise ExcelInvalido("Use cuatro columnas y como máximo 500 filas de usuarios.")
        iterator = sheet.iter_rows(max_row=MAX_ROWS + 1, max_col=4)
        header = next(iterator)
        if tuple(str(cell.value or "").strip().lower() for cell in header) != HEADERS:
            raise ExcelInvalido("Encabezados requeridos: id_externo, nombre, email, rol.")
        rows = []
        for number, cells in enumerate(iterator, start=2):
            values = [cell.value for cell in cells]
            if all(value is None or value == "" for value in values):
                continue
            formula = any(cell.data_type == "f" for cell in cells)
            raw = {}
            invalid_type = False
            for key, cell in zip(HEADERS, cells):
                value = cell.value
                if cell.data_type == "f":
                    value = ""
                if key == "id_externo" and type(value) is int and abs(value) < 10 ** 15:
                    value = str(value)
                if value is None:
                    value = ""
                if not isinstance(value, str):
                    invalid_type = True
                    value = ""
                raw[key] = value.strip()
            raw["nombre"] = " ".join(raw["nombre"].split())
            raw["email"] = raw["email"].lower()
            raw["rol"] = raw["rol"].upper()
            data = {key: value[:300] for key, value in raw.items()}
            row = {"fila": number, "datos": data, "estado": "PENDIENTE",
                   "codigo": "VALIDA", "mensaje": "Lista para importar.", "usuario_id": None}
            if formula:
                row.update(estado="INVALIDA", codigo="FORMULA", mensaje="No se admiten fórmulas en las celdas de usuarios.")
            elif invalid_type or raw["rol"] not in ROLES or len(raw["email"]) > 150:
                row.update(estado="INVALIDA", codigo="DATOS_INVALIDOS",
                           mensaje="Use texto y uno de los tres roles de líder de la plantilla.")
            else:
                try:
                    identity = IdentidadCrear(**raw, compania_id=compania_id)
                    if len(str(identity.email)) > 150:
                        row.update(estado="INVALIDA", codigo="DATOS_INVALIDOS", mensaje="El correo supera 150 caracteres.")
                    else:
                        row["datos"] = {key: str(getattr(identity, key)) for key in HEADERS}
                except ValidationError:
                    row.update(estado="INVALIDA", codigo="DATOS_INVALIDOS",
                               mensaje="Revise identificación (1-50), nombre (1-100), correo y rol.")
            rows.append(row)
        if not rows:
            raise ExcelInvalido("La hoja Usuarios no contiene usuarios.")
        ids = Counter(row["datos"]["id_externo"] for row in rows if row["datos"]["id_externo"])
        emails = Counter(row["datos"]["email"] for row in rows if row["datos"]["email"])
        for row in rows:
            if row["estado"] == "PENDIENTE" and (
                ids[row["datos"]["id_externo"]] > 1 or emails[row["datos"]["email"]] > 1
            ):
                row.update(estado="INVALIDA", codigo="DUPLICADA_ARCHIVO",
                           mensaje="Identificación o correo repetido en este archivo.")
        return rows
    except ExcelInvalido:
        raise
    except Exception as error:
        raise ExcelInvalido("No se pudo leer el Excel. Utilice una copia de la plantilla sin fórmulas.") from error
    finally:
        if workbook is not None:
            workbook.close()

def _estilo(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B67")
    for column in ("A", "B", "C", "D", "E", "F", "G"):
        sheet.column_dimensions[column].width = 30

def _texto(cell, value):
    cell.value = str(value if value is not None else "")
    cell.data_type = "s"
    cell.number_format = "@"

def plantilla_excel() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Usuarios"
    sheet.append(HEADERS)
    for row in sheet.iter_rows(min_row=2, max_row=501, max_col=4):
        for cell in row:
            cell.number_format = "@"
    roles = DataValidation(type="list", formula1='"' + ",".join(ROLES) + '"', allow_blank=False)
    roles.errorTitle = "Rol no permitido"
    roles.error = "Seleccione un rol de líder."
    roles.showErrorMessage = True
    sheet.add_data_validation(roles)
    roles.add("D2:D501")
    _estilo(sheet)
    guide = workbook.create_sheet("Guia")
    for line in [
        "Complete Usuarios con hasta 500 filas. Seleccione la compañía en IAM, no en el archivo.",
        "No agregue contraseñas, tokens, columnas adicionales, fórmulas o macros.",
        "Para conservar ceros iniciales, introduzca la identificación como texto.",
        "Roles permitidos: " + ", ".join(ROLES),
        "El correo debe pertenecer al nuevo usuario; allí recibirá su contraseña temporal.",
        "La creación requiere SMTP operativo. Cada usuario cambia contraseña y configura su propio MFA.",
        "Cada líder recibe habilitadas las tres vistas de su área; soporte puede ajustar casos puntuales en la grilla.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 115
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()

def resultado_excel(rows: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Resultados"
    sheet.append(("fila", *HEADERS, "estado", "resultado"))
    for row in rows:
        values = [row["fila"], *(row["datos"].get(key, "") for key in HEADERS),
                  row["estado"], row["mensaje"]]
        for index, value in enumerate(values, start=1):
            _texto(sheet.cell(row=sheet.max_row + 1 if index == 1 else sheet.max_row, column=index), value)
    _estilo(sheet)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
