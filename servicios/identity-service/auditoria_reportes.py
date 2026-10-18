"""Informe XLSX de auditoría persistida, reservado a ADMIN_TI."""
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
import json
from uuid import uuid4

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError

from auditoria import _sanitizar_valor
from database import engine
from seguridad import requiere_roles, SESSION_INACTIVITY_MINUTES

router = APIRouter(prefix="/auditoria", tags=["Informes de auditoría"])
admin_actor = requiere_roles("ADMIN_TI")
MAX_DAYS = 31
MAX_ROWS = 5000
MAX_DETAIL = 2000
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

CONSULTA = text("""
    SELECT e.id, e.evento_uuid, e.fecha_evento, e.servicio, e.evento, e.resultado,
           e.actor_usuario_id, e.objetivo_usuario_id, e.compania_id, e.detalle,
           u.nombre AS actor_nombre, u.email AS actor_email
    FROM auditoria_eventos e
    LEFT JOIN usuarios_plataforma u ON u.id = e.actor_usuario_id
    WHERE e.fecha_evento >= :inicio AND e.fecha_evento < :fin
      AND (:evento IS NULL OR e.evento = :evento)
      AND (:servicio IS NULL OR e.servicio = :servicio)
      AND (:resultado IS NULL OR e.resultado = :resultado)
      AND (:usuario_id IS NULL OR e.actor_usuario_id = :usuario_id
           OR e.objetivo_usuario_id = :usuario_id)
      AND (:compania_id IS NULL OR e.compania_id = :compania_id)
    ORDER BY e.fecha_evento, e.id
    LIMIT :limite
""")


def admin_vigente(connection, actor):
    # La autorización vigente y la lectura comparten transacción; no se confía
    # en el menú ni en roles enviados por el navegador.
    valid = connection.execute(text("""
        SELECT u.id FROM usuarios_plataforma u
        JOIN credenciales cr ON cr.usuario_id = u.id
        JOIN sesiones s ON s.usuario_id = u.id
        JOIN mfa_totp m ON m.usuario_id = u.id
        JOIN usuario_roles_plataforma ur ON ur.usuario_id = u.id
        JOIN roles_plataforma r ON r.id = ur.rol_id
        WHERE u.id = :uid AND s.id = :sid AND u.activo
          AND cr.activo AND NOT cr.bloqueado AND NOT cr.requiere_cambio_password
          AND (cr.password_expira_en IS NULL OR cr.password_expira_en > NOW())
          AND m.activo AND NOT s.revocada AND s.fecha_expiracion > NOW()
          AND COALESCE(s.ultima_actividad, s.fecha_inicio) >
              NOW() - (:idle * INTERVAL '1 minute')
          AND r.activo AND r.nombre = 'ADMIN_TI'
        FOR SHARE OF u, cr, s, m, ur, r
    """), {"uid": actor["id"], "sid": actor.get("session_id"),
           "idle": SESSION_INACTIVITY_MINUTES}).first()
    if not valid:
        raise HTTPException(403, "Se requiere una sesión ADMIN TI vigente con MFA.")


def texto_excel(value):
    value = "" if value is None else str(value)
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def agregar_fila(sheet, values):
    sheet.append([texto_excel(value) for value in values])
    for cell in sheet[sheet.max_row]:
        # Nunca interpretar los datos auditados como fórmulas de Excel.
        cell.data_type = "s"
        cell.number_format = "@"


def informe_excel(rows, filters, actor_id, generated):
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = "Auditoria"
        agregar_fila(sheet, ["ID evento", "UUID evento", "Fecha UTC", "Servicio",
                            "Evento", "Resultado", "ID usuario actor",
                            "Nombre usuario actor (actual)", "Usuario actor (correo actual)",
                            "ID usuario objetivo", "ID compañía", "Detalle (resumen)"])
        for row in rows:
            detail = json.dumps(_sanitizar_valor(row["detalle"]), ensure_ascii=False,
                                separators=(",", ":"))
            if len(detail) > MAX_DETAIL:
                detail = detail[:MAX_DETAIL] + "… [detalle resumido]"
            actor_name = row["actor_nombre"] or (
                "Sin usuario identificado" if row["actor_usuario_id"] is None
                else "Usuario no disponible")
            agregar_fila(sheet, [row["id"], row["evento_uuid"],
                                row["fecha_evento"].astimezone(timezone.utc).isoformat(),
                                row["servicio"], row["evento"], row["resultado"],
                                row["actor_usuario_id"], actor_name, row["actor_email"],
                                row["objetivo_usuario_id"],
                                row["compania_id"], detail])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B67")
        for column in ("A", "G", "J", "K"):
            sheet.column_dimensions[column].width = 22
        for column in ("B", "C", "D", "E", "F"):
            sheet.column_dimensions[column].width = 34
        for column in ("H", "I"):
            sheet.column_dimensions[column].width = 44
        sheet.column_dimensions["L"].width = 75
        summary = workbook.create_sheet("Resumen")
        for values in [
            ["Informe", "Auditoría de usuarios y accesos"],
            ["Generado UTC", generated.isoformat()], ["Generado por ID", actor_id],
            ["Desde UTC (incluido)", filters["desde"]],
            ["Hasta UTC (incluido)", filters["hasta"]],
            ["Eventos incluidos", len(rows)],
            ["Evento", filters["evento"] or "Todos"],
            ["Servicio", filters["servicio"] or "Todos"],
            ["Resultado", filters["resultado"] or "Todos"],
            ["Usuario (actor u objetivo)", filters["usuario_id"] or "Todos"],
            ["Compañía", filters["compania_id"] or "Todas"],
            ["Origen", "Eventos persistidos en auditoria_eventos por Audit Service."],
            ["Identificación del actor", "Nombre y correo actuales del usuario asociado al ID histórico; no son una copia de sus datos al ocurrir el evento."],
            ["Entrega pendiente", "Eventos aún pendientes del dispatcher no se incluyen."],
            ["Detalle", "Sanitizado y resumido a 2000 caracteres; no contiene credenciales."],
            ["Límites", "31 días y 5000 eventos por descarga; no se omiten filas por el límite."],
        ]:
            agregar_fila(summary, values)
        summary.column_dimensions["A"].width = 32
        summary.column_dimensions["B"].width = 100
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


@router.get("/informe.xlsx")
def descargar_informe(
    desde: date,
    hasta: date,
    evento: str | None = Query(None, min_length=1, max_length=100, pattern=r"^[A-Z][A-Z0-9_]*$"),
    servicio: str | None = Query(None, min_length=1, max_length=100, pattern=r"^[a-z][a-z0-9_.-]*$"),
    resultado: str | None = Query(None, min_length=1, max_length=50, pattern=r"^[A-Z][A-Z0-9_]*$"),
    usuario_id: int | None = Query(None, gt=0, le=2147483647),
    compania_id: int | None = Query(None, gt=0, le=2147483647),
    actor=Depends(admin_actor),
):
    days = (hasta - desde).days + 1
    if not 1 <= days <= MAX_DAYS:
        raise HTTPException(422, "Seleccione un período de 1 a 31 días, con fecha final igual o posterior a la inicial.")
    if hasta == date.max:
        raise HTTPException(422, "La fecha final está fuera del rango admitido.")
    filters = {"desde": desde.isoformat(), "hasta": hasta.isoformat(),
               "evento": evento, "servicio": servicio, "resultado": resultado,
               "usuario_id": usuario_id, "compania_id": compania_id}
    generated = datetime.now(timezone.utc)
    try:
        with engine.begin() as connection:
            admin_vigente(connection, actor)
            parameters = {**filters, "inicio": datetime.combine(desde, time.min, timezone.utc),
                          "fin": datetime.combine(hasta + timedelta(days=1), time.min, timezone.utc),
                          "limite": MAX_ROWS + 1}
            rows = connection.execute(CONSULTA, parameters).mappings().all()
            if len(rows) > MAX_ROWS:
                raise HTTPException(422, "El informe supera 5000 eventos. Reduzca las fechas o aplique más filtros.")
            content = informe_excel(rows, filters, actor["id"], generated)
            connection.execute(text("""
                INSERT INTO auditoria_outbox(evento_uuid, fecha_evento, servicio,
                    evento, resultado, actor_usuario_id, detalle)
                VALUES (:uuid, :fecha, 'identity', 'AUDITORIA_INFORME_GENERADO',
                        'EXITO', :actor, CAST(:detalle AS JSONB))
            """), {"uuid": uuid4(), "fecha": generated, "actor": actor["id"],
                   "detalle": json.dumps({"filtros": filters, "eventos": len(rows), "formato": "xlsx"})})
    except SQLAlchemyError:
        raise HTTPException(503, "No se pudo generar el informe. Verifique la disponibilidad de auditoría.") from None
    return Response(content, media_type=XLSX_TYPE, headers={
        "Content-Disposition": f'attachment; filename="auditoria-{desde.isoformat()}-{hasta.isoformat()}.xlsx"',
        "Cache-Control": "no-store", "Pragma": "no-cache",
    })
