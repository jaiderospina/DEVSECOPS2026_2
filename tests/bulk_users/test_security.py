import pytest
from sqlalchemy import text
from cargas_limite import MAX_REQUEST_BYTES

@pytest.mark.parametrize("uid",[2,4])
def test_leader_and_admin_ti_cannot_bulk_create(client,headers,uid):
    assert client.get("/usuarios/cargas/plantilla",headers=headers(uid)).status_code==403

def test_unauthenticated_template_denied(client):
    assert client.get("/usuarios/cargas/plantilla").status_code in (401,403)

def test_other_company_denied(client,headers,excel):
    r=client.post("/usuarios/cargas/validar",headers=headers(),data={"compania_id":"2"},files={"archivo":("a.xlsx",excel())})
    assert r.status_code==403

@pytest.mark.parametrize("query",[
    "UPDATE mfa_totp SET activo=FALSE WHERE usuario_id=1",
    "UPDATE credenciales SET requiere_cambio_password=TRUE WHERE usuario_id=1",
    "UPDATE credenciales SET bloqueado=TRUE WHERE usuario_id=1",
    "UPDATE sesiones SET revocada=TRUE,fecha_revocacion=NOW() WHERE usuario_id=1",
    "UPDATE sesiones SET ultima_actividad=NOW()-INTERVAL '1 day' WHERE usuario_id=1",
    "DELETE FROM usuario_companias_plataforma WHERE usuario_id=1",
    "UPDATE companias SET activo=FALSE WHERE id=1",
])
def test_changed_authorization_blocks_existing_batch(preview,client,db,query):
    h,job=preview()
    with db.begin() as c:c.execute(text(query))
    r=client.get("/usuarios/cargas/"+job["id"],headers=h)
    assert r.status_code in (401,403)

def test_foreign_actor_cannot_read_confirm_create_or_download(preview,client,headers):
    _,job=preview();h=headers(3);prefix="/usuarios/cargas/"+job["id"]
    assert client.get(prefix,headers=h).status_code==404
    assert client.get(prefix+"/resultado",headers=h).status_code==404
    assert client.post(prefix+"/confirmar",headers=h,json={"confirmar":True}).status_code==404
    assert client.post(prefix+"/filas/2",headers=h).status_code==404

def test_confirmation_expires_and_retention_blocks_access(preview,client,db):
    h,job=preview()
    with db.begin() as c:c.execute(text("UPDATE cargas_usuarios SET confirmar_antes=NOW()-INTERVAL '1 minute'"))
    assert client.post("/usuarios/cargas/"+job["id"]+"/confirmar",headers=h,json={"confirmar":True}).status_code==409
    with db.begin() as c:c.execute(text("UPDATE cargas_usuarios SET conservar_hasta=NOW()-INTERVAL '1 minute'"))
    assert client.get("/usuarios/cargas/"+job["id"],headers=h).status_code==404

def test_rate_limit_and_list_scoped_to_actor(preview,client,headers):
    h=headers()
    for _ in range(10):preview(auth=h)
    r=client.get("/usuarios/cargas",headers=h,params={"compania_id":1})
    assert len(r.json()["cargas"])==10
    other=client.get("/usuarios/cargas",headers=headers(3),params={"compania_id":2})
    assert other.json()["cargas"]==[]
    # Reutilizar una plantilla real para el undécimo lote.
    workbook=client.get("/usuarios/cargas/plantilla",headers=h)
    from io import BytesIO
    from openpyxl import load_workbook
    w=load_workbook(BytesIO(workbook.content));w["Usuarios"].append(["1","Persona","p@example.org","LIDER_FACTURACION"])
    # La plantilla ya contiene 500 celdas formateadas: escribir la fila dos.
    for n,value in enumerate(["1","Persona","p@example.org","LIDER_FACTURACION"],1):w["Usuarios"].cell(2,n,value)
    w["Usuarios"].delete_rows(502);out=BytesIO();w.save(out);w.close()
    assert client.post("/usuarios/cargas/validar",headers=h,data={"compania_id":"1"},files={"archivo":("a.xlsx",out.getvalue())}).status_code==429

def test_request_size_rejected_before_upload_parse(client,headers):
    r=client.post("/usuarios/cargas/validar",headers={**headers(),"Content-Length":str(MAX_REQUEST_BYTES+1)},content=b"x")
    assert r.status_code==413

def test_invalid_row_never_creates_account(preview,client,db):
    h,job=preview([("1","Persona","p@example.org","ADMIN_TI"),("2","Otra","otra@example.org","LIDER_FACTURACION")])
    client.post("/usuarios/cargas/"+job["id"]+"/confirmar",headers=h,json={"confirmar":True})
    result=client.post("/usuarios/cargas/"+job["id"]+"/filas/2",headers=h)
    assert result.json()["fila"]["estado"]=="INVALIDA"
    with db.connect() as c:assert c.execute(text("SELECT count(*) FROM usuarios_plataforma")).scalar_one()==4
