"""Autorización real JWT/PostgreSQL y contenido seguro del informe de auditoría."""
from datetime import datetime, timezone
from io import BytesIO
import json
import secrets
from uuid import uuid4

from openpyxl import load_workbook
import pytest
from sqlalchemy import text
import auditoria_reportes as reports

URL = "/auditoria/informe.xlsx"
PERIOD = {"desde": "2026-10-01", "hasta": "2026-10-07"}


@pytest.fixture(autouse=True)
def clean_audit(db):
    # Solo la base efímera del job: las tablas auditadas de producción son append-only.
    with db.begin() as connection:
        connection.execute(text("TRUNCATE auditoria_eventos RESTART IDENTITY"))


def event(db, **changes):
    values = {"uuid": uuid4(), "fecha": datetime(2026, 10, 4, 12, tzinfo=timezone.utc),
              "servicio": "identity", "evento": "LOGIN_EXITOSO", "resultado": "EXITO",
              "actor": 1, "objetivo": 2, "compania": 1, "detalle": {"motivo": "Acceso de prueba"}, **changes}
    values["detalle"] = json.dumps(values["detalle"])
    with db.begin() as connection:
        return connection.execute(text("""
            INSERT INTO auditoria_eventos(evento_uuid,fecha_evento,servicio,evento,
                 resultado,actor_usuario_id,objetivo_usuario_id,compania_id,detalle)
            VALUES (:uuid,:fecha,:servicio,:evento,:resultado,:actor,:objetivo,
                    :compania,CAST(:detalle AS JSONB)) RETURNING id
        """), values).scalar_one()


def workbook(response):
    assert response.status_code == 200, response.text if response.status_code != 200 else ""
    return load_workbook(BytesIO(response.content), data_only=False)


def test_admin_downloads_excel_and_generation_is_audited(client, headers, db):
    event(db)
    response = client.get(URL, params=PERIOD, headers=headers(4))
    book = workbook(response)
    try:
        assert book.sheetnames == ["Auditoria", "Resumen"]
        assert book["Auditoria"].max_row == 2
        assert book["Auditoria"]["E2"].value == "LOGIN_EXITOSO"
        assert book["Auditoria"].auto_filter.ref == "A1:L2"
        assert book["Auditoria"]["G2"].value == "1"
        assert book["Auditoria"]["H2"].value == "Usuario 1"
        assert book["Auditoria"]["I2"].value == "usuario1@example.org"
        assert book["Auditoria"]["J2"].value == "2"
        summary = dict(book["Resumen"].values)
        assert summary["Eventos incluidos"] == "1"
        assert summary["Generado por ID"] == "4"
        assert summary["Desde UTC (incluido)"] == PERIOD["desde"]
    finally:
        book.close()
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["content-type"].startswith(reports.XLSX_TYPE)
    assert response.headers["content-disposition"] == 'attachment; filename="auditoria-2026-10-01-2026-10-07.xlsx"'
    with db.connect() as connection:
        audit = connection.execute(text("SELECT actor_usuario_id,detalle,estado FROM auditoria_outbox WHERE evento='AUDITORIA_INFORME_GENERADO'")).mappings().one()
        assert audit["actor_usuario_id"] == 4 and audit["estado"] == "PENDIENTE"
        assert audit["detalle"]["eventos"] == 1
        assert audit["detalle"]["filtros"]["hasta"] == PERIOD["hasta"]


@pytest.mark.parametrize("uid", [1, 2, 3])
def test_other_roles_cannot_download(client, headers, uid):
    assert client.get(URL, params=PERIOD, headers=headers(uid)).status_code == 403


def test_anonymous_cannot_download(client):
    assert client.get(URL, params=PERIOD).status_code == 401


@pytest.mark.parametrize("change", ["role", "mfa", "session", "temporary", "blocked"])
def test_revoked_or_restricted_admin_is_denied(client, headers, db, change):
    auth = headers(4)
    mutations = {
        "role": "DELETE FROM usuario_roles_plataforma WHERE usuario_id=4",
        "mfa": "UPDATE mfa_totp SET activo=FALSE WHERE usuario_id=4",
        "session": "UPDATE sesiones SET revocada=TRUE, fecha_revocacion=NOW() WHERE usuario_id=4",
        "temporary": "UPDATE credenciales SET requiere_cambio_password=TRUE WHERE usuario_id=4",
        "blocked": "UPDATE credenciales SET bloqueado=TRUE WHERE usuario_id=4",
    }
    with db.begin() as connection:
        connection.execute(text(mutations[change]))
    assert client.get(URL, params=PERIOD, headers=auth).status_code in (401, 403)


def test_inclusive_utc_days_and_stable_event_order(client, headers, db):
    event(db, fecha=datetime(2026, 9, 30, 23, 59, 59, tzinfo=timezone.utc))
    first = event(db, fecha=datetime(2026, 10, 1, tzinfo=timezone.utc))
    last = event(db, fecha=datetime(2026, 10, 7, 23, 59, 59, tzinfo=timezone.utc))
    event(db, fecha=datetime(2026, 10, 8, tzinfo=timezone.utc))
    book = workbook(client.get(URL, params=PERIOD, headers=headers(4)))
    try:
        assert [row[0] for row in list(book["Auditoria"].values)[1:]] == [str(first), str(last)]
    finally:
        book.close()


@pytest.mark.parametrize("filters", [
    {"evento":"LOGIN_EXITOSO"}, {"servicio":"identity"}, {"resultado":"EXITO"},
    {"usuario_id":2}, {"usuario_id":1}, {"compania_id":1},
    {"evento":"LOGIN_EXITOSO", "servicio":"identity", "usuario_id":2, "compania_id":1},
])
def test_each_filter_and_combination(client, headers, db, filters):
    event(db)
    event(db, servicio="permissions", evento="ACCESO_DENEGADO", resultado="DENEGADO", actor=3, objetivo=3, compania=2)
    book = workbook(client.get(URL, params={**PERIOD, **filters}, headers=headers(4)))
    try:
        assert book["Auditoria"].max_row == 2
        assert book["Auditoria"]["E2"].value == "LOGIN_EXITOSO"
    finally:
        book.close()


def test_empty_report_keeps_headers_and_summary(client, headers):
    book = workbook(client.get(URL, params=PERIOD, headers=headers(4)))
    try:
        assert book["Auditoria"].max_row == 1
        assert dict(book["Resumen"].values)["Eventos incluidos"] == "0"
    finally:
        book.close()


@pytest.mark.parametrize("invalid", [
    {"desde":"2026-10-08"}, {"hasta":"2026-11-01"}, {"desde":"2026-02-30"},
    {"desde":"9999-12-31", "hasta":"9999-12-31"}, {"usuario_id":0},
    {"usuario_id":2147483648}, {"compania_id":-1}, {"evento":"LOGIN' OR 1=1 --"},
    {"servicio":"https://untrusted.invalid"},
])
def test_invalid_filters_are_rejected(client, headers, invalid):
    assert client.get(URL, params={**PERIOD, **invalid}, headers=headers(4)).status_code == 422


def test_excess_rows_are_rejected_without_partial_download(client, headers, db, monkeypatch):
    monkeypatch.setattr(reports, "MAX_ROWS", 2)
    for _ in range(3):
        event(db)
    response = client.get(URL, params=PERIOD, headers=headers(4))
    assert response.status_code == 422 and "5000" in response.json()["detail"]
    with db.connect() as connection:
        assert connection.execute(text("SELECT count(*) FROM auditoria_outbox WHERE evento='AUDITORIA_INFORME_GENERADO'")).scalar_one() == 0


def test_formula_text_controls_and_nested_secrets_are_safe(client, headers, db):
    secret = secrets.token_hex(16)
    event(db, servicio="=2+3", evento="@SUM(1)", resultado="-1+2", detalle={
        "password":secret, "lista":[{"smtp_password":secret,"access_token":secret}],
        "motivo":"Texto\x01 sin control", "detalle_largo":"x" * 2200})
    book = workbook(client.get(URL, params=PERIOD, headers=headers(4)))
    try:
        sheet = book["Auditoria"]
        assert sheet["D2"].value == "=2+3" and sheet["D2"].data_type == "s"
        assert all(cell.data_type != "f" for row in sheet for cell in row)
        detail = sheet["L2"].value
        assert secret not in detail and "[REDACTED]" in detail
        assert "\x01" not in detail and "[detalle resumido]" in detail
        assert len(detail) < 2050
    finally:
        book.close()


def test_database_failure_does_not_expose_internal_details(client, headers, monkeypatch):
    from sqlalchemy.exc import SQLAlchemyError
    class UnavailableDatabase:
        def begin(self):
            raise SQLAlchemyError("internal-configuration-must-not-be-returned")
    monkeypatch.setattr(reports, "engine", UnavailableDatabase())
    response = client.get(URL, params=PERIOD, headers=headers(4))
    assert response.status_code == 503
    assert "internal-configuration" not in response.text


@pytest.mark.parametrize("actor,label", [
    (None, "Sin usuario identificado"), (9876, "Usuario no disponible"),
])
def test_missing_actor_does_not_drop_or_misattribute_events(client, headers, db, actor, label):
    event(db, actor=actor)
    book = workbook(client.get(URL, params=PERIOD, headers=headers(4)))
    try:
        sheet = book["Auditoria"]
        assert sheet.max_row == 2
        assert sheet["G2"].value == (None if actor is None else str(actor))
        assert sheet["H2"].value == label
        assert sheet["I2"].value is None
        assert sheet["J2"].value == "2"
    finally:
        book.close()


def test_actor_uses_current_directory_data_including_inactive_users(client, headers, db):
    event(db)
    with db.begin() as connection:
        connection.execute(text("UPDATE usuarios_plataforma SET nombre=:name, email=:email, activo=FALSE WHERE id=1"),
                           {"name": "Nombre actualizado", "email": "actual@example.org"})
    book = workbook(client.get(URL, params={**PERIOD, "usuario_id": 1}, headers=headers(4)))
    try:
        sheet = book["Auditoria"]
        assert sheet["H2"].value == "Nombre actualizado"
        assert sheet["I2"].value == "actual@example.org"
        assert "actual" in sheet["H1"].value and "actual" in sheet["I1"].value
        assert "no son una copia" in dict(book["Resumen"].values)["Identificación del actor"]
    finally:
        book.close()


def test_actor_directory_fields_remain_safe_excel_text(client, headers, db):
    with db.begin() as connection:
        connection.execute(text("UPDATE usuarios_plataforma SET nombre=:name, email=:email WHERE id=1"),
                           {"name": '=HYPERLINK("https://example.org")\x01', "email": "=1+1@example.org"})
    event(db)
    book = workbook(client.get(URL, params=PERIOD, headers=headers(4)))
    try:
        sheet = book["Auditoria"]
        assert sheet["H2"].value == '=HYPERLINK("https://example.org")'
        assert sheet["I2"].value == "=1+1@example.org"
        assert sheet["H2"].data_type == sheet["I2"].data_type == "s"
        assert all(cell.data_type != "f" for row in sheet for cell in row)
    finally:
        book.close()
