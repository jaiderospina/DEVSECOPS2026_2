from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from openpyxl import load_workbook
from cargas_excel import ExcelInvalido, leer_usuarios, plantilla_excel, resultado_excel, MAX_FILE_BYTES

def test_template_and_leading_zero_preservation(excel):
    w=load_workbook(BytesIO(plantilla_excel()))
    assert w.sheetnames==["Usuarios","Guia"]
    assert w["Usuarios"]["A2"].number_format=="@"
    assert w["Usuarios"].data_validations.count==1
    w.close()
    row=leer_usuarios(excel(), "usuarios.xlsx",1)[0]
    assert row["datos"]["id_externo"]=="001" and row["estado"]=="PENDIENTE"

@pytest.mark.parametrize("role",["ADMIN_TI","ADMIN_SOPORTE","DESCONOCIDO"])
def test_forbidden_roles_are_not_importable(excel,role):
    row=leer_usuarios(excel([("1","Persona","p@example.org",role)]),"a.xlsx",1)[0]
    assert row["estado"]=="INVALIDA"

@pytest.mark.parametrize("column,value",[(0,"=1+1"),(1,"=HYPERLINK(\"https://example.org\")"),(2,"=1+1")])
def test_formula_rejected_and_not_retained(excel,column,value):
    data=["001","Persona","p@example.org","LIDER_FACTURACION"];data[column]=value
    row=leer_usuarios(excel([data]),"a.xlsx",1)[0]
    assert row["codigo"]=="FORMULA"
    assert value not in str(row)

@pytest.mark.parametrize("duplicate",["id","email"])
def test_all_duplicate_rows_are_flagged(excel,duplicate):
    rows=[["001","Una Persona","uno@example.org","LIDER_FACTURACION"],
          ["001" if duplicate=="id" else "002","Otra Persona","UNO@example.org" if duplicate=="email" else "dos@example.org","LIDER_LOGISTICA"]]
    assert all(r["codigo"]=="DUPLICADA_ARCHIVO" for r in leer_usuarios(excel(rows),"a.xlsx",1))

@pytest.mark.parametrize("filename",["a.xls","a.xlsm","a.csv"])
def test_non_xlsx_rejected(excel,filename):
    with pytest.raises(ExcelInvalido): leer_usuarios(excel(),filename,1)

def test_header_sheet_and_row_limits(excel):
    for content in (excel(headers=["nombre","email","rol","id_externo"]),excel(title="Otra"),
                    excel([("1","Persona","uno@example.org","LIDER_FACTURACION")]*501)):
        with pytest.raises(ExcelInvalido): leer_usuarios(content,"a.xlsx",1)

def test_500_rows_accepted(excel):
    rows=[(str(n),"Persona "+str(n),"u"+str(n)+"@example.org","LIDER_LOGISTICA") for n in range(500)]
    assert len(leer_usuarios(excel(rows),"a.xlsx",1))==500

def test_nonzip_oversize_and_compression_bomb_rejected():
    for content in (b"not-a-workbook",b"x"*(MAX_FILE_BYTES+1)):
        with pytest.raises(ExcelInvalido): leer_usuarios(content,"a.xlsx",1)
    out=BytesIO()
    with ZipFile(out,"w",ZIP_DEFLATED) as z: z.writestr("xl/worksheets/sheet1.xml","x"*1000000)
    with pytest.raises(ExcelInvalido): leer_usuarios(out.getvalue(),"a.xlsx",1)

@pytest.mark.parametrize("entry",["xl/vbaProject.bin","xl/externalLinks/externalLink1.xml"])
def test_embedded_macro_or_external_link_rejected(excel,entry):
    out=BytesIO(excel())
    with ZipFile(out,"a") as z:z.writestr(entry,b"fixture")
    with pytest.raises(ExcelInvalido):leer_usuarios(out.getvalue(),"a.xlsx",1)

def test_report_untrusted_values_are_text_not_formulas():
    row={"fila":2,"datos":{"id_externo":"=1+1","nombre":"+CMD","email":"@text","rol":"LIDER_LOGISTICA"},"estado":"ERROR","mensaje":"=HYPERLINK"}
    w=load_workbook(BytesIO(resultado_excel([row])))
    assert all(c.data_type=="s" for c in w["Resultados"][2])
    assert w["Resultados"]["B2"].value=="=1+1"
    w.close()

def test_physical_rows_and_columns_cannot_bypass_forged_dimensions(excel):
    import re
    for rows in ([("1","Persona","p@example.org","LIDER_FACTURACION","extra")],
                 [("1","Persona","p@example.org","LIDER_FACTURACION")]*501):
        original=BytesIO(excel(rows));out=BytesIO()
        with ZipFile(original) as src,ZipFile(out,"w",ZIP_DEFLATED) as dst:
            for entry in src.infolist():
                content=src.read(entry.filename)
                if entry.filename=="xl/worksheets/sheet1.xml":
                    content=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:D2"',content)
                dst.writestr(entry,content)
        with pytest.raises(ExcelInvalido):leer_usuarios(out.getvalue(),"a.xlsx",1)
